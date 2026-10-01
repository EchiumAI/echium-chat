"""Server-side text extraction for uploaded workspace files.

Agents can only use a document if it has a text body (``text_s3_key``). Chat
attachments reach the model as native Bedrock document blocks, but files
uploaded to Files (and attachments persisted as documents) only had their raw
bytes stored — so sharing them with an agent gave it nothing to read. This
module turns the common office/text formats into plain text at upload time.

Best-effort by design: a failure returns "" and the upload still succeeds.
"""

from __future__ import annotations

import csv
import io
import logging
import re
from html import unescape

logger = logging.getLogger(__name__)

# Keep extracted text bounded so one huge file can't bloat S3/embeddings.
MAX_TEXT_CHARS = 1_000_000

_TEXT_EXTENSIONS = {
    "txt",
    "md",
    "markdown",
    "json",
    "yaml",
    "yml",
    "xml",
    "log",
    "ini",
    "toml",
    "py",
    "js",
    "ts",
    "tsx",
    "jsx",
    "java",
    "go",
    "rs",
    "rb",
    "sh",
    "sql",
    "excalidraw",
}


def _ext(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _decode(body: bytes) -> str:
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return body.decode(enc)
        except UnicodeDecodeError:
            continue
    return body.decode("utf-8", errors="replace")


def _from_html(body: bytes) -> str:
    text = _decode(body)
    text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = unescape(text)
    return re.sub(r"[ \t]+", " ", text)


def _from_csv(body: bytes) -> str:
    reader = csv.reader(io.StringIO(_decode(body)))
    return "\n".join("\t".join(row) for row in reader)


def _from_pdf(body: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(body))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - one bad page shouldn't lose the rest
            logger.debug("pypdf failed on a page", exc_info=True)
    return "\n\n".join(p for p in pages if p.strip())


def _from_docx(body: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(body))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _from_xlsx(body: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"## {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def body_to_plain_text(text: str, content_type: str = "") -> str:
    """Normalise a stored document body for the model.

    The Docs editor saves HTML (``text/html``); agents should see readable
    text, not markup, and embeddings should not be built from tags. Markdown
    and plain text pass through unchanged.
    """
    if not text:
        return ""
    ctype = (content_type or "").lower()
    looks_like_html = "html" in ctype or (
        text.lstrip().startswith("<") and re.search(r"</(p|div|h[1-6]|li|ul|ol)>", text)
    )
    if not looks_like_html:
        return text
    try:
        return _from_html(text.encode("utf-8")).strip()
    except Exception:  # noqa: BLE001
        return text


def extract_text(body: bytes, filename: str, content_type: str = "") -> str:
    """Return plain text for a file, or "" if the format isn't supported."""
    if not body:
        return ""
    ext = _ext(filename)
    ctype = (content_type or "").lower()
    try:
        if ext == "pdf" or ctype == "application/pdf":
            text = _from_pdf(body)
        elif ext == "docx" or "wordprocessingml" in ctype:
            text = _from_docx(body)
        elif ext in ("xlsx", "xlsm") or "spreadsheetml" in ctype:
            text = _from_xlsx(body)
        elif ext == "csv" or ctype == "text/csv":
            text = _from_csv(body)
        elif ext in ("html", "htm") or ctype == "text/html":
            text = _from_html(body)
        elif ext in _TEXT_EXTENSIONS or ctype.startswith("text/"):
            text = _decode(body)
        else:
            # Unknown binary (images, zips, legacy .doc/.xls): nothing to read.
            return ""
    except Exception:  # noqa: BLE001 - extraction must never break an upload
        logger.warning(f"Text extraction failed for {filename}", exc_info=True)
        return ""
    text = text.strip()
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS]
    return text
